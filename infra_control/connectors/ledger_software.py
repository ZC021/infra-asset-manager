from __future__ import annotations

# Software license ledger connector.
# 소프트웨어 관리대장(구매/보유 라이선스)을 software_ledger 테이블로 적재한다.
# - 읽기전용 소스(엑셀). 라이선스 키 등 시크릿은 마스킹해서만 저장한다.
# - norm_name 으로 SWeeper software_inventory_summary 와 대조(컴플라이언스)에 쓴다.
import os
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from infra_control.db import connect, encode_json, init_db

try:
    import openpyxl  # .17/로컬 venv 에 존재
except Exception:  # pragma: no cover
    openpyxl = None

ROOT = Path(__file__).resolve().parents[2]
SOURCE_ID = "software-ledger"

DEFAULT_LEDGER_FILES = tuple(
    Path(d) / "ACME_소프트웨어_구매_정리.xlsx"
    for d in (str(Path.home() / "Downloads"), "/opt/infra-asset-manager/var/imports/ledger", str(ROOT / "var" / "imports" / "ledger"))
)

# 시트 -> (제품 헤더명, 기본 라이선스 유형)
SHEETS = {
    "전체 요약": ("소프트웨어", "summary"),
    "라이선스 키": ("소프트웨어", "key"),
    "비쥬얼데이타 발주 상세": ("제품", "order"),
    "Microsoft Open Value": ("제품", "volume"),
    "직접결제 정기구독": ("서비스", "saas"),
}

NOTE_TERMS = ("범례", "총 라이선스", "인증번호", "발행번호", "관리 콘솔", "즉시 사용",
              "별도 키 없음", "제품번호 확인", "위 시트", "welcome letter", "계약 시작")
KEY_HEADERS = ("값", "key", "라이선스 키", "시리얼", "serial", "인증")
GROUP_SUFFIX = ("통합그룹", "패키지 그룹", "패키지그룹")
ALIASES = {
    "한글": "hancom hwp", "한컴오피스 한글": "hancom hwp", "한컴오피스": "hancom office",
    "hancom office hwp": "hancom hwp",
    "비즈니스용 microsoft 365 앱": "microsoft 365",
    "microsoft 365 apps for business": "microsoft 365",
    "microsoft 365 for business": "microsoft 365",
}


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def norm(s: str) -> str:
    s = (s or "").lower().strip()
    for g in GROUP_SUFFIX:
        s = s.replace(g.lower(), "")
    s = re.sub(r"[()\[\]_/.,#:\-]+", " ", s)
    s = re.sub(r"\bv?\d+(\.\d+)*\b", " ", s)
    s = s.replace("professional", "pro").replace("standard", "std")
    s = re.sub(r"\s+", " ", s).strip()
    return ALIASES.get(s, s)


def mask_key(value: str) -> str:
    v = (value or "").strip()
    if not v:
        return ""
    compact = re.sub(r"\s+", "", v)
    if len(compact) <= 8:
        return "****"
    return f"{compact[:4]}****{compact[-4:]}"


def last_date(text: str) -> str:
    # "계약 ... (2026-01-15 ~ 2029-01-31)" 등에서 마지막 날짜를 만료로
    dates = re.findall(r"\d{4}[-./]\d{1,2}[-./]\d{1,2}", text or "")
    return dates[-1].replace(".", "-").replace("/", "-") if dates else ""


def license_type(sheet_default: str, *texts: str) -> str:
    # 깔끔한 taxonomy: perpetual / subscription / volume / unknown
    blob = " ".join(t for t in texts if t).lower()
    if sheet_default == "saas" or any(k in blob for k in ("구독", "subscription", "annual", "정기", "월 정산", "saas")):
        return "subscription"
    if sheet_default == "volume" or any(k in blob for k in ("open value", "vlk", "volume")):
        return "volume"
    if any(k in blob for k in ("영구", "perpetual", "copy", "누적", "open")):
        return "perpetual"
    return "unknown"


def _header_index(header_row: list[str], *keywords: str) -> int | None:
    for i, h in enumerate(header_row):
        hl = (h or "").lower()
        if any(k.lower() in hl for k in keywords):
            return i
    return None


def parse_ledger(path: Path) -> list[dict[str, Any]]:
    if openpyxl is None:
        raise RuntimeError("openpyxl required to parse the software ledger")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    items: list[dict[str, Any]] = []
    for sheet, (prod_col_name, sheet_default) in SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        rows = [[("" if c is None else str(c).strip()) for c in r] for r in ws.iter_rows(values_only=True)]
        # 헤더 행 찾기
        hidx = None
        for i, r in enumerate(rows):
            if prod_col_name in r:
                hidx = i
                break
        if hidx is None:
            continue
        hdr = rows[hidx]
        pc = hdr.index(prod_col_name)
        idx = {
            "vendor": _header_index(hdr, "구분", "공급", "구매처", "벤더"),
            "supply": _header_index(hdr, "공급", "구매처"),
            "qty": _header_index(hdr, "수량"),
            "amount": _header_index(hdr, "금액"),
            "ident": _header_index(hdr, "식별자 유형"),
            "key": _header_index(hdr, *KEY_HEADERS),
            "activated": _header_index(hdr, "발급", "활성일", "일자", "결제일", "주문 확인일"),
            "contract": _header_index(hdr, "계약", "결제 형태", "보증 기간", "플랜"),
            "note": _header_index(hdr, "비고", "상태"),
        }

        def cell(row: list[str], j: int | None) -> str:
            return row[j] if (j is not None and j < len(row)) else ""

        for r in rows[hidx + 1:]:
            if len(r) <= pc:
                continue
            product = r[pc]
            if not product or product == "-":
                continue
            if any(t in product for t in NOTE_TERMS) or len(product) > 60:
                continue
            contract = cell(r, idx["contract"])
            vendor = cell(r, idx["vendor"])
            supply = cell(r, idx["supply"])
            note = cell(r, idx["note"])
            keyval = cell(r, idx["key"])
            items.append({
                "product": product,
                "norm_name": norm(product),
                "vendor": vendor,
                "source_sheet": sheet,
                "license_type": license_type(sheet_default, contract, vendor, supply, note, product),
                "qty": cell(r, idx["qty"]),
                "has_key": 1 if keyval and keyval not in ("-", "team") else 0,
                "key_masked": mask_key(keyval) if keyval and keyval not in ("-", "team") else "",
                "identifier_type": cell(r, idx["ident"]),
                "activated_at": cell(r, idx["activated"]),
                "expiry": last_date(contract),
                "amount_usd": cell(r, idx["amount"]),
                "note": note,
            })
    return items


def _resolve_path(explicit: str | None) -> Path | None:
    if explicit:
        p = Path(explicit)
        return p if p.exists() else None
    env = os.environ.get("INFRA_CONTROL_SW_LEDGER")
    if env and Path(env).exists():
        return Path(env)
    for cand in DEFAULT_LEDGER_FILES:
        if cand.exists():
            return cand
    return None


def import_software_ledger(path: str | None = None, db_path: Path | None = None) -> dict[str, Any]:
    ledger = _resolve_path(path)
    init_db(db_path) if db_path else init_db()
    if ledger is None:
        return {"status": "skipped", "reason": "ledger file not found", "product_rows": 0}
    try:
        # 암호화(예: 보유대장.xlsx CDFV2)/손상 파일이어도 bootstrap 전체를 죽이지 않는다.
        items = parse_ledger(ledger)
    except Exception as exc:
        return {"status": "error", "reason": f"{type(exc).__name__}: {exc}",
                "ledger_file": str(ledger), "product_rows": 0}
    imported_at = utc_now()
    target = db_path
    cm = connect(target) if target else connect()
    with cm as conn:
        conn.execute("DELETE FROM software_ledger")
        conn.executemany(
            """
            INSERT INTO software_ledger (
              product, norm_name, vendor, source_sheet, license_type, qty,
              has_key, key_masked, identifier_type, activated_at, expiry,
              amount_usd, note, collected_at, imported_at
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            """,
            [
                (
                    it["product"], it["norm_name"], it["vendor"], it["source_sheet"],
                    it["license_type"], it["qty"], it["has_key"], it["key_masked"],
                    it["identifier_type"], it["activated_at"], it["expiry"],
                    it["amount_usd"], it["note"], imported_at, imported_at,
                )
                for it in items
            ],
        )
        conn.execute(
            """
            INSERT OR REPLACE INTO integration_sources (
              id, name, kind, endpoint, status, last_sync_at, metadata_json
            ) VALUES (?,?,?,?,?,?,?)
            """,
            (SOURCE_ID, "Software License Ledger", "software-ledger-source",
             str(ledger), "read-only", imported_at,
             encode_json({"sheets": list(SHEETS.keys())})),
        )
        conn.execute(
            """
            INSERT INTO software_ledger_runs (source_id, imported_at, ledger_file, product_rows, detail_json)
            VALUES (?,?,?,?,?)
            """,
            (SOURCE_ID, imported_at, str(ledger), len(items),
             encode_json({"saas": sum(1 for i in items if i["license_type"] == "subscription"),
                          "with_key": sum(1 for i in items if i["has_key"])})),
        )
    return {"status": "ok", "product_rows": len(items), "ledger_file": str(ledger),
            "saas": sum(1 for i in items if i["license_type"] == "subscription")}


def main() -> int:
    import argparse
    import json
    from infra_control.db import DB_PATH

    parser = argparse.ArgumentParser()
    parser.add_argument("--db", default=str(DB_PATH))
    parser.add_argument("--ledger", default=None)
    args = parser.parse_args()
    result = import_software_ledger(path=args.ledger, db_path=Path(args.db))
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
